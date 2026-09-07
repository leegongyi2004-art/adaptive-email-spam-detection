"""Build a submission-ready Word (.docx) from REPORT_FYP2.md.

Produces REPORT_FYP2.docx formatted to UTAR FYP house style:
  * Times New Roman; body 12 pt with 1.5 line spacing, justified
  * black TNR chapter/section headings (chapter titles on a new page, centred)
  * centred cover / title pages
  * real Word tables with bold, shaded header rows
  * ready-made figures embedded automatically where the PNG exists
     (run reports/make_figures.py and reports/make_result_figures.py first)
  * page numbers in the footer; [FILL IN] / [FIGURE ...] placeholders left visible

Usage (from the project root):
    python reports/build_docx.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
MD = ROOT / "REPORT_FYP2.md"
OUT = ROOT / "REPORT_FYP2.docx"
FIGS = ROOT / "reports" / "figures"

FIG_MAP = {
    "FIGURE 1.1": "fig1_1_spam_share.png",
    "FIGURE 1.2": "fig1_2_phishing_attacks.png",
    "FIGURE 1.3": "fig1_3_threat_stats.png",
    "FIGURE 2.1": "fig2_1_taxonomy.png",
    "FIGURE 2.2": "fig2_2_rule_based.png",
    "FIGURE 2.3": "fig2_3_content_statistical.png",
    "FIGURE 2.4": "fig2_4_metadata.png",
    "FIGURE 2.5": "fig2_5_hybrid_fusion.png",
    "FIGURE 3.1": "fig3_1_system_architecture.png",
    "FIGURE 3.2": "fig3_2_use_case.png",
    "FIGURE 3.3": "fig3_3_activity.png",
    "FIGURE 4.1": "fig4_1_block_diagram.png",
    "FIGURE 6.1": "fig6_1_confusion_matrix.png",
    "FIGURE 6.2": "fig6_2_roc.png",
    "FIGURE 6.3": "fig6_3_threshold_catch.png",
    "FIGURE 6.4": "fig6_4_adaptive_roc.png",
}

CENTER_PAGES = {"FRONT COVER", "TITLE PAGE"}
TNR = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")


def add_runs(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            par.add_run(part)


def set_font(style, size, bold=False, italic=False, color=BLACK):
    style.font.name = TNR
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), TNR)


def style_doc(doc):
    # Page margins: left 1.2 in (binding/cover), right/top/bottom 1 in
    for sec in doc.sections:
        sec.left_margin = Inches(1.2)
        sec.right_margin = Inches(1.0)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)

    normal = doc.styles["Normal"]
    set_font(normal, 12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    set_font(doc.styles["Heading 1"], 16, bold=True)
    set_font(doc.styles["Heading 2"], 14, bold=True)
    set_font(doc.styles["Heading 3"], 12, bold=True)
    set_font(doc.styles["Heading 4"], 12, bold=True, italic=True)
    for nm in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        doc.styles[nm].paragraph_format.space_before = Pt(12)
        doc.styles[nm].paragraph_format.space_after = Pt(8)


def add_page_numbers(doc):
    for sec in doc.sections:
        footer = sec.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
        fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
        run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
        run.font.name = TNR; run.font.size = Pt(10)


def shade(cell, hexfill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), hexfill)
    tcpr.append(shd)


def caption(par):
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in par.runs:
        r.font.size = Pt(11); r.bold = True
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(8)


def add_image(doc, marker):
    key = marker[1:marker.index(":")].strip().upper()
    fname = FIG_MAP.get(key)
    if fname and (FIGS / fname).exists():
        doc.add_picture(str(FIGS / fname), width=Inches(6.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = marker.split(":", 1)[1].strip().rstrip("].")
        p = doc.add_paragraph(); add_runs(p, key.title() + " — " + cap); caption(p)
        return
    p = doc.add_paragraph()
    r = p.add_run(marker.strip("[]")); r.italic = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
        i += 1
    rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
    return rows, i


def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, r in enumerate(rows):
        for ci in range(ncol):
            cell = table.cell(ri, ci)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            par.paragraph_format.space_after = Pt(2)
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(par, r[ci] if ci < len(r) else "")
            for run in par.runs:
                run.font.name = TNR; run.font.size = Pt(10.5)
                if ri == 0:
                    run.bold = True
            if ri == 0:
                shade(cell, "D9E2F3")
    doc.add_paragraph()


def main():
    lines = MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_doc(doc)
    add_page_numbers(doc)

    buf: list[str] = []
    in_center = False
    in_code = False
    skip_comment = False
    list_re = re.compile(r"^(\s*)([-*]|\d+\.)\s+")
    n = len(lines)

    def flush_prose():
        nonlocal buf
        if not buf:
            return
        text = " ".join(x.strip() for x in buf).strip()
        if text:
            p = doc.add_paragraph()
            add_runs(p, text)
            if in_center:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(10)
                for r in p.runs:
                    r.bold = True
        buf = []

    i = 0
    while i < n:
        raw = lines[i]
        s = raw.strip()

        if s.startswith("<!--"):
            flush_prose()
            skip_comment = ("-->" not in s)
            i += 1; continue
        if skip_comment:
            if "-->" in s:
                skip_comment = False
            i += 1; continue

        if s.startswith("```"):
            flush_prose()
            in_code = not in_code; i += 1; continue
        if in_code:
            p = doc.add_paragraph()
            r = p.add_run(raw if raw else " ")
            r.font.name = "Consolas"; r.font.size = Pt(9.5)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after = Pt(0)
            i += 1; continue

        if s.startswith("|"):
            flush_prose(); rows, ni = parse_table(lines, i); add_table(doc, rows); i = ni; continue

        if s.startswith("[FIGURE"):
            flush_prose()
            marker = s
            while "]" not in marker and i + 1 < n:
                i += 1; marker += " " + lines[i].strip()
            add_image(doc, marker); i += 1; continue

        if s.startswith(">") or "Ready-made image" in s:
            flush_prose(); i += 1; continue
        if s == "---":
            flush_prose(); i += 1; continue

        if s.startswith("#"):
            flush_prose()
            m = re.match(r"^(#+)\s+(.*)", s)
            level = len(m.group(1)); text = m.group(2).strip()
            if level == 1:
                doc.add_page_break()
                h = doc.add_heading(level=1); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                in_center = text in CENTER_PAGES
            else:
                h = doc.add_heading(level=min(level - 1, 4))
                h.alignment = WD_ALIGN_PARAGRAPH.LEFT
                in_center = False
            run = h.add_run(text); run.font.name = TNR; run.font.color.rgb = BLACK
            i += 1; continue

        if not s:
            flush_prose(); i += 1; continue

        if s.startswith("**Table"):
            flush_prose(); p = doc.add_paragraph(); add_runs(p, s); caption(p); i += 1; continue

        # A bullet/number list marker starts a NEW paragraph; wrapped lines below
        # it (and ordinary prose) accumulate into the current paragraph.
        if list_re.match(s):
            flush_prose()
        buf.append(s)
        i += 1

    flush_prose()
    doc.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
