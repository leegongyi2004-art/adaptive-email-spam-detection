#!/usr/bin/env python3
"""
make_changes_docx.py
--------------------
Builds reports/paste/CHANGES_ONLY.docx - a clean Word document containing ONLY the
sections that changed in the latest polishing round. Unlike the .txt version:
  * headings are real Word headings,
  * tables are real, pre-formatted Word tables (copy -> paste arrives as a table),
  * figure-insert points and instructions are clearly shaded.
Workflow for the student: open CHANGES_ONLY.docx, and for each banner block, go to
that section in the real report, delete the old text/table, and copy the content
under the banner across.
"""

import os
import re
import sys

import build_docx as B          # reuse styling + table helpers (guarded main)
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

from make_changes_sheet import BLOCKS, slice_lines, references_block

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "REPORT_FYP2.md")
OUT = os.path.join(HERE, "paste", "CHANGES_ONLY.docx")

DARK = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x55, 0x55, 0x55)
YELLOW = "FFF2CC"   # figure-insert shading
BLUE = "D9E2F3"     # banner shading


def strip_inline(text):
    """Drop backticks; keep **bold**/*italic* for add_runs to render."""
    return text.replace("`", "")


def banner(doc, n, total, title):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    B.shade(cell, BLUE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"REPLACE IN YOUR REPORT  ({n} of {total})")
    r.bold = True
    r.font.size = Pt(11)
    r.font.name = B.TNR
    r.font.color.rgb = DARK
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(title)
    r2.font.size = Pt(10)
    r2.font.name = B.TNR
    r2.font.color.rgb = GREY
    p3 = cell.add_paragraph()
    p3.paragraph_format.space_after = Pt(0)
    r3 = p3.add_run("Go to this section in your own report, delete its old contents, "
                    "then copy everything under this banner (tables copy as tables).")
    r3.italic = True
    r3.font.size = Pt(9)
    r3.font.name = B.TNR
    r3.font.color.rgb = GREY


def figure_box(doc, num, caption):
    from docx.shared import Inches
    key = f"FIGURE {num}"
    fname = B.FIG_MAP.get(key)
    path = (B.FIGS / fname) if fname else None
    if path and path.exists():
        # ready-made image: embed it so it copies as a picture
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(path), width=Inches(5.8))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        rc = cap.add_run(f"Figure {num}: {caption}")
        rc.italic = True
        rc.font.size = Pt(9)
        rc.font.name = B.TNR
        rc.font.color.rgb = GREY
    else:
        # user-supplied figure (screenshots / their-PC charts): yellow placeholder
        table = doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        cell = table.cell(0, 0)
        B.shade(cell, YELLOW)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"INSERT YOUR OWN PICTURE HERE  \u2014  Figure {num}: {caption}")
        r.bold = True
        r.font.size = Pt(10)
        r.font.name = B.TNR
        r.font.color.rgb = RGBColor(0x80, 0x60, 0x00)


def code_block(doc, code_lines):
    for j, ln in enumerate(code_lines):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0 if j < len(code_lines) - 1 else 6)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(ln if ln else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)


def table_no_trailing_blank(doc, rows):
    from docx.enum.table import WD_TABLE_ALIGNMENT
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, r in enumerate(rows):
        for ci in range(ncol):
            cell = table.cell(ri, ci)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.line_spacing = 1.0
            par.paragraph_format.space_after = Pt(2)
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT
            B.add_runs(par, r[ci] if ci < len(r) else "")
            for run in par.runs:
                run.font.name = B.TNR
                run.font.size = Pt(10.5)
                if ri == 0:
                    run.bold = True
            if ri == 0:
                B.shade(cell, "D9E2F3")


def parse_tokens(lines):
    """Turn a slice of markdown lines into (kind, payload) tokens."""
    tokens = []
    buf = []
    in_comment = False

    def flush():
        if buf:
            txt = " ".join(buf).strip()
            if txt:
                tokens.append(("p", strip_inline(txt)))
            buf.clear()

    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if "<!--" in s:
            in_comment = True
        if in_comment:
            if "-->" in s:
                in_comment = False
            i += 1
            continue
        if not s:
            flush()
        elif s.startswith(">"):
            flush()  # drop build-note blockquotes
        elif s.startswith("#"):
            flush()
            level = len(s) - len(s.lstrip("#"))
            tokens.append(("h", (level, strip_inline(s.lstrip("#").strip()))))
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c or "") for c in cells):
                    rows.append([strip_inline(c) for c in cells])
                i += 1
            tokens.append(("table", rows))
            continue
        elif s.startswith("EQ#"):
            flush()
            tokens.append(("eq", strip_inline(s[3:].strip())))
        elif s.startswith("[FIGURE"):
            flush()
            joined = s
            while "]" not in joined and i + 1 < len(lines):
                i += 1
                joined += " " + lines[i].strip()
            m = re.match(r"\[FIGURE\s*([0-9]+\.[0-9]+)\s*:?\s*(.*?)\]", joined)
            num = m.group(1) if m else ""
            cap = m.group(2) if m else joined
            tokens.append(("figure", (num, strip_inline(cap.rstrip(".")))))
        elif s.startswith("```"):
            flush()
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            tokens.append(("code", code))
        else:
            buf.append(s)
        i += 1
    flush()
    return tokens


def render_tokens(doc, tokens):
    for kind, payload in tokens:
        if kind == "h":
            level, text = payload
            h = doc.add_heading(text, level=min(level, 4))
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(6)
        elif kind == "p":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
            pf.space_after = Pt(6)
            pf.space_before = Pt(0)
            B.add_runs(p, payload)
        elif kind == "table":
            table_no_trailing_blank(doc, payload)
        elif kind == "eq":
            pe = doc.add_paragraph()
            pe.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pe.paragraph_format.space_before = Pt(4)
            pe.paragraph_format.space_after = Pt(4)
            re_ = pe.add_run(payload)
            re_.font.name = "Cambria Math"
            re_.font.size = Pt(12)
        elif kind == "figure":
            figure_box(doc, payload[0], payload[1])
        elif kind == "code":
            code_block(doc, payload)


def main():
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = B.Document()
    B.style_doc(doc)

    # Title / how-to
    h = doc.add_heading("CHANGES ONLY \u2014 paste-ready sections", level=0)
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.LEFT
    B.add_runs(intro,
        "This document contains ONLY the sections that changed in the latest round, in order. "
        "For each blue banner, open the matching section of your own report, delete the old "
        "text/table, and copy the heading, paragraphs and table from here. Tables are already "
        "real tables \u2014 they paste as tables. Yellow boxes mark where a picture goes. "
        "Anything not shown here is unchanged. Do not include the blue banner or yellow boxes "
        "in your final report.")

    total = len(BLOCKS)
    for n, (title, start, end) in enumerate(BLOCKS, 1):
        if n > 1:
            doc.add_page_break()
        banner(doc, n, total, title)
        if start is None:
            for entry in references_block(lines):
                if entry.strip():
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    pf = p.paragraph_format
                    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
                    pf.space_after = Pt(6)
                    B.add_runs(p, strip_inline(entry))
        else:
            render_tokens(doc, parse_tokens(slice_lines(lines, start, end)))

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
